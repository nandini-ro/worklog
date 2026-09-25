import io

from openpyxl import load_workbook

from tests.conftest import category_id, make_entry, make_project

RANGE = {"start_date": "2026-09-21", "end_date": "2026-09-25"}


def _seed(client, auth):
    ap = make_project(client, auth, name="Agent Platform")
    mcp = make_project(client, auth, name="MCP Integration")
    garak = make_project(client, auth, name="Garak Testing")
    dev, test = category_id(client, auth, "Development"), category_id(client, auth, "Testing")
    make_entry(client, auth, work_date="2026-09-25", project_id=ap["id"], category_id=dev,
               task_title="Model API configuration", description="Added model API configuration through the frontend.")
    make_entry(client, auth, work_date="2026-09-25", project_id=mcp["id"], category_id=test,
               task_title="MCP connectivity testing", description="Tested MCP server connectivity.")
    make_entry(client, auth, work_date="2026-09-24", project_id=garak["id"], category_id=test, status="in_progress",
               task_title="Garak endpoint prep", description="Prepared agent endpoints for security testing.")
    make_entry(client, auth, work_date="2026-09-23", project_id=mcp["id"], status="blocked",
               task_title="Staging deploy", blockers="Waiting for staging credentials")
    make_entry(client, auth, work_date="2026-09-10", project_id=ap["id"], task_title="Out of range")
    return ap


def test_report_structure(client, auth):
    _seed(client, auth)
    r = client.get("/api/reports", params=RANGE, headers=auth)
    assert r.status_code == 200
    rep = r.json()
    assert rep["period"]["label"] == "21 Sep 2026 – 25 Sep 2026"
    assert rep["user"]["name"] == "Alice"
    assert rep["totals"] == {"entries": 4, "completed": 2, "in_progress": 1, "planned": 0, "blocked": 1,
                             "projects": 3, "active_days": 3}
    assert [d["date"] for d in rep["daily"]] == ["2026-09-25", "2026-09-24", "2026-09-23"]
    assert [g["project"] for g in rep["daily"][0]["groups"]] == ["Agent Platform", "MCP Integration"]
    assert [e["task_title"] for e in rep["in_progress"]] == ["Garak endpoint prep"]
    assert rep["blockers"][0]["blockers"] == "Waiting for staging credentials"
    assert {p["project"] for p in rep["project_summary"]} == {"Agent Platform", "MCP Integration", "Garak Testing"}
    assert "Agent Platform" in rep["summary"] and rep["summary_source"] == "auto"
    assert "Out of range" not in str(rep)
    for word in ("hours", "duration"):
        assert word not in str(rep).lower()


def test_report_filters(client, auth):
    ap = _seed(client, auth)
    rep = client.get("/api/reports", params={**RANGE, "project_id": ap["id"]}, headers=auth).json()
    assert rep["totals"]["entries"] == 1 and rep["filters"]["project"] == "Agent Platform"
    rep = client.get("/api/reports", params={**RANGE, "status": "blocked"}, headers=auth).json()
    assert [e["task_title"] for e in rep["blockers"]] == ["Staging deploy"] and rep["totals"]["entries"] == 1
    rep = client.get("/api/reports", params={**RANGE, "category_id": category_id(client, auth, "Testing")},
                     headers=auth).json()
    assert rep["totals"]["entries"] == 2


def test_report_invalid_range_and_empty(client, auth):
    r = client.get("/api/reports", params={"start_date": "2026-09-25", "end_date": "2026-09-01"}, headers=auth)
    assert r.status_code == 422
    rep = client.get("/api/reports", params=RANGE, headers=auth).json()
    assert rep["totals"]["entries"] == 0 and "No work was recorded" in rep["summary"]


def test_custom_summary_used(client, auth):
    _seed(client, auth)
    rep = client.get("/api/reports", params={**RANGE, "summary": "My custom summary."}, headers=auth).json()
    assert rep["summary"] == "My custom summary." and rep["summary_source"] == "custom"


def test_export_pdf(client, auth):
    _seed(client, auth)
    r = client.get("/api/reports/export/pdf", params=RANGE, headers=auth)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 2000
    assert 'filename="work-report-alice-2026-09-21_to_2026-09-25.pdf"' in r.headers["content-disposition"]


def test_export_pdf_empty_period(client, auth):
    r = client.get("/api/reports/export/pdf", params=RANGE, headers=auth)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_export_xlsx(client, auth):
    _seed(client, auth)
    r = client.get("/api/reports/export/xlsx", params=RANGE, headers=auth)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Daily Work", "Projects", "Blockers"]
    rows = list(wb["Daily Work"].iter_rows(values_only=True))
    assert rows[0] == ("Date", "Project", "Category", "Task", "Description", "Status", "Notes", "Blockers")
    assert len(rows) == 5
    assert wb["Blockers"]["E2"].value == "Waiting for staging credentials"


def test_export_csv(client, auth):
    _seed(client, auth)
    r = client.get("/api/reports/export/csv", params=RANGE, headers=auth)
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    text = r.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0] == "Date,Project,Category,Task,Description,Status,Notes,Blockers"
    assert len(lines) == 5 and "Model API configuration" in text
    assert "hours" not in lines[0].lower()


def test_exports_require_auth(client):
    for fmt in ("pdf", "xlsx", "csv"):
        assert client.get(f"/api/reports/export/{fmt}", params=RANGE).status_code == 401
