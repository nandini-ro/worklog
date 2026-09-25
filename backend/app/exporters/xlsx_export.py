import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.exporters.csv_export import HEADERS, _rows

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font = HEADER_FILL, HEADER_FONT


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_xlsx(report: dict) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = report["title"]
    ws["A1"].font = Font(bold=True, size=16)
    meta = [
        ("Name", report["user"]["name"]),
        ("Period", report["period"]["label"]),
        ("Generated", report["generated_label"]),
    ]
    for k in ("project", "category", "status"):
        if report["filters"][k]:
            meta.append((f"Filter: {k.title()}", report["filters"][k]))
    r = 3
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Summary").font = Font(bold=True, size=12)
    r += 1
    ws.cell(row=r, column=1, value=report["summary"]).alignment = WRAP
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = max(30, 15 * (len(report["summary"]) // 90 + 1))
    r += 2
    t = report["totals"]
    for k, v in [("Work items", t["entries"]), ("Completed", t["completed"]), ("In progress", t["in_progress"]),
                 ("Planned", t["planned"]), ("Blocked", t["blocked"]), ("Projects", t["projects"]),
                 ("Active days", t["active_days"])]:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    _widths(ws, [22, 40, 30, 30])

    ws = wb.create_sheet("Daily Work")
    _header(ws, 1, HEADERS)
    for i, row in enumerate(_rows(report), start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v).alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(1, ws.max_row)}"
    _widths(ws, [12, 22, 16, 36, 60, 13, 36, 36])

    ws = wb.create_sheet("Projects")
    _header(ws, 1, ["Project", "Work items", "Completed", "In Progress", "Planned", "Blocked", "Tasks"])
    for i, p in enumerate(report["project_summary"], start=2):
        vals = [p["project"], p["total"], p["completed"], p["in_progress"], p["planned"], p["blocked"],
                "\n".join(f"• {t}" for t in p["tasks"])]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(ws, [24, 11, 11, 12, 10, 10, 70])

    ws = wb.create_sheet("Blockers")
    _header(ws, 1, ["Date", "Project", "Task", "Status", "Blocker"])
    for i, e in enumerate(report["blockers"], start=2):
        for j, v in enumerate([e["date"], e["project"] or "", e["task_title"], e["status_label"], e["blockers"] or ""], start=1):
            ws.cell(row=i, column=j, value=v).alignment = WRAP
    _widths(ws, [12, 22, 40, 13, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
